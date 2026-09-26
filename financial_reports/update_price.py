"""Fill the Price row of the comps workbook with a closing price from Yahoo Finance.

Takes the close on --date, or the latest close before it if the market was shut
that day (weekend / holiday). Default date is today.

Usage:
    python financial_reports/update_price.py ORCL Oracle templates/TradingComps.xlsx --date 2026-09-25
"""

import argparse
from datetime import date, datetime, timedelta, timezone

import requests
from openpyxl import load_workbook

from update_comps import check_formulas, resize_notes, write_cell

CHART_URL = "https://query1.finance.yahoo.com/v8/finance/chart/{ticker}"
PRICE_ROW = 19
PRICE_FMT = '#,##0.00;(#,##0.00);"-"'


def closing_price(ticker, on_or_before):
    start = datetime.combine(on_or_before - timedelta(days=10), datetime.min.time(), timezone.utc)
    end = datetime.combine(on_or_before + timedelta(days=1), datetime.min.time(), timezone.utc)
    resp = requests.get(
        CHART_URL.format(ticker=ticker),
        params={"period1": int(start.timestamp()), "period2": int(end.timestamp()), "interval": "1d"},
        headers={"User-Agent": "Mozilla/5.0"},
        timeout=30,
    )
    resp.raise_for_status()
    result = resp.json()["chart"]["result"][0]
    offset = result["meta"]["gmtoffset"]
    closes = [
        (datetime.fromtimestamp(ts + offset, timezone.utc).date(), close)
        for ts, close in zip(result["timestamp"], result["indicators"]["quote"][0]["close"])
        if close is not None
    ]
    closes = [(d, c) for d, c in closes if d <= on_or_before]
    if not closes:
        raise SystemExit(f"No closing price for {ticker} on or before {on_or_before}")
    day, close = closes[-1]
    return day, round(close, 2), result["meta"]["currency"]


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("ticker", help="Yahoo Finance ticker, e.g. ORCL")
    ap.add_argument("company", help="Company name as in row 1 of the workbook")
    ap.add_argument("workbook")
    ap.add_argument("--date", type=date.fromisoformat, default=date.today())
    ap.add_argument("--sheet")
    args = ap.parse_args()

    day, price, currency = closing_price(args.ticker, args.date)

    wb = load_workbook(args.workbook)
    ws = wb[args.sheet] if args.sheet else wb.active
    col = next(
        (c for c in range(2, ws.max_column + 1)
         if str(ws.cell(row=1, column=c).value or "").strip().lower() == args.company.lower()),
        None,
    )
    if col is None:
        raise SystemExit(f"{args.company!r} not found in row 1; populate the EV bridge first")

    note = f"Yahoo Finance\n{args.ticker} closing price ({currency}), {day.strftime('%d %b %Y')}\nfinance.yahoo.com/quote/{args.ticker}/history"
    write_cell(ws, PRICE_ROW, col, price, note, PRICE_FMT)
    resize_notes(ws)
    wb.calculation.fullCalcOnLoad = True
    wb.save(args.workbook)
    print(f"{args.company}: {args.ticker} close {price} {currency} on {day}")
    if not check_formulas(args.workbook):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
