"""Write extracted EV-bridge line items into the Trading Comps workbook.

Input is a JSON file produced from reading the filings, e.g.
``filings/ORCL/ev_bridge.json``:

    {
      "company": "Oracle",
      "currency": "USD",
      "items": {
        "4":  {"value": "=-10786-1003", "note": "Oracle QR Aug 26\\nCash and cash equivalents\\npg 5\\nMarketable securities\\npg 5"},
        "13": {"value": 0, "note": "Oracle QR Aug 26\\nPreferred stock - nil per balance sheet\\npg 5"}
      }
    }

Every item must carry a note: each cell is written with its value, Aptos 10 blue
font, the number format and a legacy yellow Note in one step. Notes are titled
"Claude Analyst" and sized so the full text shows on hover. The Price row and
the formula rows are never touched.

The company goes into the column whose row 1 already holds its name, otherwise
into the first empty column from B onwards.

Usage:
    python financial_reports/update_comps.py templates/TradingComps.xlsx filings/ORCL/ev_bridge.json
"""

import argparse
import json
import math
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font

BLUE = "000000FF"
NOTE_AUTHOR = "Claude Analyst"
# Excel draws notes in ~9pt Tahoma: about 7px per character and 15px per line.
NOTE_CHAR_PX, NOTE_LINE_PX, NOTE_PAD_PX = 7, 15, 24
NOTE_MIN_WIDTH, NOTE_MAX_WIDTH = 180, 440
DEFAULT_NUM_FMT = '#,##0.0;(#,##0.0);"-"'
WRITABLE_ROWS = {4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 16}
PROTECTED_ROWS = {18, 19, 20, 22, 23, 24, 25}
ERROR_TOKENS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!")


def target_column(ws, company):
    for col in range(2, ws.max_column + 2):
        name = ws.cell(row=1, column=col).value
        if name and str(name).strip().lower() == company.lower():
            return col
    col = 2
    while ws.cell(row=1, column=col).value not in (None, ""):
        col += 1
    return col


def existing_number_format(ws, col):
    """Use the format already live in the sheet (pre-flight check), not an assumed one."""
    for c in [col] + list(range(2, ws.max_column + 1)):
        fmt = ws.cell(row=4, column=c).number_format
        if fmt and fmt != "General":
            return fmt
    return DEFAULT_NUM_FMT


def note_size(text):
    """Width/height in px so the whole note is visible on hover, without scrolling."""
    lines = text.split("\n")
    longest = max(len(line) for line in lines)
    width = min(max(longest * NOTE_CHAR_PX + NOTE_PAD_PX, NOTE_MIN_WIDTH), NOTE_MAX_WIDTH)
    chars_per_line = (width - NOTE_PAD_PX) // NOTE_CHAR_PX
    wrapped = sum(max(1, math.ceil(len(line) / chars_per_line)) for line in lines)
    return width, wrapped * NOTE_LINE_PX + NOTE_PAD_PX


def write_cell(ws, row, col, value, note_text, num_fmt):
    cell = ws.cell(row=row, column=col)
    cell.value = value
    cell.font = Font(name="Aptos", size=10, color=BLUE)
    cell.number_format = num_fmt
    text = f"{NOTE_AUTHOR}:\n{note_text}"
    note = Comment(text, NOTE_AUTHOR)
    note.width, note.height = note_size(text)
    cell.comment = note


def resize_notes(ws):
    """openpyxl forgets note sizes when it reloads a file, so re-apply them before every save."""
    for row in ws.iter_rows():
        for cell in row:
            if cell.comment:
                cell.comment.width, cell.comment.height = note_size(cell.comment.text)


NOTE_RUN_PROPS = '<sz val="9"/><color indexed="81"/><rFont val="Tahoma"/><family val="2"/>'


def bold_note_titles(path):
    """Make the "Claude Analyst:" title line of every note bold.

    openpyxl only writes plain-text notes, so after saving rewrite each note's
    text in xl/comments*.xml as a bold title run plus a regular body run, which is
    how Excel itself stores the author line of a note.
    """
    title = re.escape(NOTE_AUTHOR + ":")
    pattern = re.compile(rf"<text><t(?: [^>]*)?>({title})(.*?)</t></text>", re.S)
    replacement = (
        f"<text><r><rPr><b/>{NOTE_RUN_PROPS}</rPr><t>\\1</t></r>"
        f'<r><rPr>{NOTE_RUN_PROPS}</rPr><t xml:space="preserve">\\2</t></r></text>'
    )
    path = Path(path)
    tmp = path.with_suffix(".tmp.xlsx")
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if re.fullmatch(r"xl/comments/comment\d+\.xml", item.filename):
                data = pattern.sub(replacement, data.decode("utf-8")).encode("utf-8")
            dst.writestr(item, data)
    tmp.replace(path)


def check_formulas(path):
    """Recalculate a copy in LibreOffice and report any Excel error values."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("warning: LibreOffice not found, skipping recalculation check", file=sys.stderr)
        return True
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(
            [soffice, "--headless", "--convert-to", "csv", "--outdir", tmp, str(path)],
            check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=120,
        )
        csv_text = (Path(tmp) / (Path(path).stem + ".csv")).read_text(errors="replace")
    errors = [t for t in ERROR_TOKENS if t in csv_text]
    if errors:
        print(f"Formula errors after recalculation: {', '.join(errors)}", file=sys.stderr)
        return False
    return True


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workbook")
    ap.add_argument("data", help="JSON file with the extracted line items")
    ap.add_argument("--sheet", help="Worksheet name (default: active sheet)")
    args = ap.parse_args()

    data = json.loads(Path(args.data).read_text())
    items = {int(k): v for k, v in data["items"].items()}

    bad = sorted(set(items) & PROTECTED_ROWS) + sorted(set(items) - WRITABLE_ROWS - PROTECTED_ROWS)
    if bad:
        sys.exit(f"Refusing to write rows {bad}: only {sorted(WRITABLE_ROWS)} are input rows")
    missing_notes = [r for r, v in items.items() if not str(v.get("note", "")).strip()]
    if missing_notes:
        sys.exit(f"Rows {missing_notes} have no source note; every written cell needs one")

    # data_only=False (default) keeps existing formulas intact.
    wb = load_workbook(args.workbook)
    ws = wb[args.sheet] if args.sheet else wb.active
    col = target_column(ws, data["company"])
    num_fmt = existing_number_format(ws, col)

    ws.cell(row=1, column=col, value=data["company"])
    if data.get("currency"):
        ws.cell(row=2, column=col, value=data["currency"])
    for row, item in sorted(items.items()):
        write_cell(ws, row, col, item["value"], item["note"], num_fmt)

    resize_notes(ws)
    wb.calculation.fullCalcOnLoad = True
    wb.save(args.workbook)
    bold_note_titles(args.workbook)
    print(f"Wrote {len(items)} cells for {data['company']} into column {ws.cell(row=1, column=col).column_letter}")
    if not check_formulas(args.workbook):
        sys.exit(1)


if __name__ == "__main__":
    main()
